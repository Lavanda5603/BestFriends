import io
from pathlib import Path
from typing import Optional
from loguru import logger

MIME_MAP = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/msword": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/plain": "txt",
    "text/csv": "txt",
}
EXT_MAP = {".pdf": "pdf", ".docx": "docx", ".doc": "docx", ".pptx": "pptx", ".xlsx": "xlsx", ".txt": "txt", ".csv": "txt"}


class DocumentService:
    def extract_text(self, data: bytes, mime_type: str, filename: str = "") -> Optional[str]:
        kind = MIME_MAP.get(mime_type) or EXT_MAP.get(Path(filename).suffix.lower())
        if not kind:
            return None
        try:
            return getattr(self, f"_extract_{kind}")(data)
        except Exception as e:
            logger.error(f"Extraction error ({kind}): {e}")
            return None

    def _extract_pdf(self, data: bytes) -> str:
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(data))
        return "\n\n".join(p.extract_text().strip() for p in r.pages if p.extract_text())

    def _extract_docx(self, data: bytes) -> str:
        from docx import Document
        return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs if p.text.strip())

    def _extract_pptx(self, data: bytes) -> str:
        from pptx import Presentation
        texts = []
        for slide in Presentation(io.BytesIO(data)).slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
        return "\n\n".join(texts)

    def _extract_xlsx(self, data: bytes) -> str:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            rows = ["\t".join(str(c) if c is not None else "" for c in row) for row in sheet.iter_rows(values_only=True)]
            parts.append(f"[{sheet.title}]\n" + "\n".join(r for r in rows if r.strip()))
        wb.close()
        return "\n\n".join(parts)

    def _extract_txt(self, data: bytes) -> str:
        for enc in ("utf-8", "cp1251", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")


document_service = DocumentService()
